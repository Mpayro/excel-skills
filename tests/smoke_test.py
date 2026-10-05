#!/usr/bin/env python3
"""Smoke test for the excel-skills scripts. Standard library plus openpyxl. No Excel needed.

  python tests/smoke_test.py

One line per check, a final count, exit 0 only when every check passed (a KNOWN DEFECT is not a failure).
Fixtures are built at run time in a temporary folder. Set SMOKE_NO_LO=1 to behave as if LibreOffice is absent.
"""
import base64, io, json, os, re, shutil, struct, subprocess, sys, tempfile, time, traceback, zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table

ROOT = Path(__file__).resolve().parent.parent
SK = ROOT / "plugins" / "excel-skills" / "skills"
PQ = SK / "excel-power-query" / "scripts"
WV = SK / "excel-write-verify" / "scripts"
PY = {p.name: p for d in (PQ, WV) for p in d.glob("*.py")}

ROWS = [("North", 10, 100, 60), ("South", 20, 250, 130), ("East", 15, 180, 90), ("West", 8, 90, 50), ("Central", 12, 140, 75)]
RES = {"ok": 0, "fail": 0, "known": 0}


def ck(name, cond, why="", known=None):
    """One check. `known` marks a script defect: it is reported, and it does not fail the run."""
    if cond:
        RES["ok"] += 1
        print("ok   " + name)
    elif known:
        RES["known"] += 1
        print(f"KNOWN DEFECT {name}: {known}")
    else:
        RES["fail"] += 1
        print(f"FAIL {name}: {why}")


def run(script, *args):
    """-> (exit code, stdout + stderr). No shell."""
    env = dict(os.environ, PYTHONUTF8="1", PYTHONIOENCODING="utf-8")
    p = subprocess.run([sys.executable, "-B", str(PY[script])] + [str(a) for a in args],
                       stdout=subprocess.PIPE, stderr=subprocess.STDOUT, env=env, timeout=600)
    return p.returncode, p.stdout.decode("utf-8", "replace").replace("\r\n", "\n")    # Windows line ends


def zip_edit(src, dst, fn):
    """Copy a zip, letting fn(parts: {name: bytes}) change it."""
    with zipfile.ZipFile(src) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    fn(parts)
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in parts.items():
            z.writestr(n, b)


def sub_part(name, pattern, repl, flags=0):
    def fn(parts):
        text = parts[name].decode("utf8")
        new, n = re.subn(pattern, repl, text, flags=flags)
        assert n, f"pattern not found in {name}: {pattern}"
        parts[name] = new.encode("utf8")
    return fn


def part(path, name):
    with zipfile.ZipFile(path) as z:
        return z.read(name).decode("utf8")


def jwrite(path, obj):
    Path(path).write_text(json.dumps(obj), encoding="utf-8")
    return path


# ---------------------------------------------------------------- fixtures

def build_base(path, summary_col="C", shared=False):
    """Data!A1:E6 = table Sales (Net = Gross - Cost), totals in G, per-row formulas in J, Summary points at a Data column."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Region", "Units", "Gross", "Cost", "Net"])
    for i, (reg, units, gross, cost) in enumerate(ROWS, start=2):
        ws.append([reg, units, gross, cost, f"=C{i}-D{i}"])
    ws.add_table(Table(displayName="Sales", ref="A1:E6"))
    ws["G1"], ws["G2"], ws["G3"] = "Total", "=SUM(E2:E6)", "=AVERAGE(C2:C6)"
    ws["J1"] = "Double"
    for r in range(2, 7):
        ws[f"J{r}"] = f"=C{r}*2"
    sm = wb.create_sheet("Summary")
    sm["A1"], sm["B1"] = "Gross total", f"=SUM(Data!{summary_col}2:{summary_col}6)"
    wb.save(path)
    tmp = Path(str(path) + ".tmp")
    # openpyxl writes an empty <v></v> in formula cells; drop it so the fixtures look like files saved by Excel.
    def drop_empty_v(parts):                 # with lxml it is <v></v>, without it <v />; either way, optional
        for name in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"):
            parts[name] = re.sub(rb"<v\s*/>|<v></v>", b"", parts[name])
    fixes = [drop_empty_v,
             sub_part("xl/tables/table1.xml", r'(<tableColumn [^>]*name="Net")\s*/>',
                      r"\1><calculatedColumnFormula>C2-D2</calculatedColumnFormula></tableColumn>")]
    if shared:       # master E2 with ref E2:E6, dependents E3:E6, each with a cached value
        def cells(m):
            r = int(m.group(1))
            f = '<f t="shared" ref="E2:E6" si="0">C2-D2</f>' if r == 2 else '<f t="shared" si="0"/>'
            return f'<c r="E{r}">{f}<v>{ROWS[r - 2][2] - ROWS[r - 2][3]}</v></c>'
        fixes.append(sub_part("xl/worksheets/sheet1.xml", r'<c r="E([2-6])"[^>]*>.*?</c>', cells, re.S))
    zip_edit(path, tmp, lambda p: [f(p) for f in fixes])
    os.replace(tmp, path)
    return path


M_CODE = '''section Section1;

shared SalesRaw = let
    Source = Excel.Workbook(File.Contents("C:\\data\\Source.xlsx"), null, true),
    DataSheet = Source{[Item="Data",Kind="Sheet"]}[Data],
    Promoted = Table.PromoteHeaders(DataSheet, [PromoteAllScalars=true])
in
    Promoted;

shared SalesNet = let
    Source = SalesRaw,
    Positive = Table.SelectRows(Source, each [Net] > 0)
in
    Positive;
'''
FILL = ('<Entry Type="FillLastUpdated" Value="d2026-01-02T03:04:05.0000000Z"/><Entry Type="FillCount" Value="l5"/>'
        '<Entry Type="FillStatus" Value="sComplete"/><Entry Type="FillErrorCount" Value="l0"/>'
        '<Entry Type="FillObjectType" Value="sTable"/>')
META = ('<LocalPackageMetadataFile><Items>'
        f'<Item><ItemLocation><ItemType>Formula</ItemType><ItemPath>Section1/SalesRaw</ItemPath></ItemLocation><StableEntries>{FILL}</StableEntries></Item>'
        '<Item><ItemLocation><ItemType>Formula</ItemType><ItemPath>Section1/SalesNet</ItemPath></ItemLocation><StableEntries/></Item>'
        '</Items></LocalPackageMetadataFile>')
MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CT = "application/vnd.openxmlformats-officedocument.spreadsheetml."


def mashup_xml():
    pkg = io.BytesIO()
    with zipfile.ZipFile(pkg, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("Formulas/Section1.m", M_CODE)
    pkg, perm, meta = pkg.getvalue(), b"<PermissionList/>", META.encode("utf-8")
    meta_blob = struct.pack("<II", 0, len(meta)) + meta
    blob = (struct.pack("<II", 0, len(pkg)) + pkg + struct.pack("<I", len(perm)) + perm
            + struct.pack("<I", len(meta_blob)) + meta_blob)
    xml = ('<?xml version="1.0" encoding="utf-16"?><DataMashup xmlns="http://schemas.microsoft.com/DataMashup">'
           + base64.b64encode(blob).decode("ascii") + "</DataMashup>")
    return xml.encode("utf-16")                                   # BOM included


def build_pq(path, tmpdir):
    """Workbook with Data!SalesRaw (4 query columns + 1 Excel-side column) fed by a hand-made mashup."""
    raw = Path(tmpdir) / "pq_raw.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Region", "Units", "Gross", "Cost", "Extra"])
    for reg, units, gross, cost in ROWS:
        ws.append([reg, units, gross, cost, "x"])
    ws.add_table(Table(displayName="SalesRaw", ref="A1:E6"))
    wb.save(raw)
    conn = (f'<connections xmlns="{MAIN}"><connection id="1" name="Query - SalesRaw" type="5" refreshedVersion="6" background="1">'
            '<dbPr connection="Provider=Microsoft.Mashup.OleDb.1;Data Source=$Workbook$;Location=SalesRaw;Extended Properties=&quot;&quot;" '
            'command="SELECT * FROM [SalesRaw]"/></connection></connections>')
    fields = "".join(f'<queryTableField id="{i}" name="{n}" tableColumnId="{i}"/>' for i, n in enumerate(["Region", "Units", "Gross", "Cost"], 1))
    qt = (f'<queryTable xmlns="{MAIN}" name="ExternalData_1" connectionId="1"><queryTableRefresh nextId="6" unboundColumnsRight="1">'
          f'<queryTableFields count="5">{fields}<queryTableField id="5" dataBound="0" tableColumnId="5"/></queryTableFields></queryTableRefresh></queryTable>')
    trel = (f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" '
            f'Type="{REL}/queryTable" Target="../queryTables/queryTable1.xml"/></Relationships>')

    def fn(p):
        p["customXml/item1.xml"] = mashup_xml()
        p["xl/connections.xml"], p["xl/queryTables/queryTable1.xml"] = conn.encode(), qt.encode()
        p["xl/tables/_rels/table1.xml.rels"] = trel.encode()
    zip_edit(raw, path, fn)
    for name, pat, rep in (
            ("[Content_Types].xml", "</Types>", f'<Override PartName="/xl/connections.xml" ContentType="{CT}connections+xml"/>'
             f'<Override PartName="/xl/queryTables/queryTable1.xml" ContentType="{CT}queryTable+xml"/></Types>'),
            ("xl/_rels/workbook.xml.rels", "</Relationships>", f'<Relationship Id="rIdC1" Type="{REL}/connections" Target="connections.xml"/></Relationships>'),
            ("xl/tables/table1.xml", "<table ", '<table tableType="queryTable" ')):
        tmp = Path(str(path) + ".tmp")
        zip_edit(path, tmp, sub_part(name, re.escape(pat), rep.replace("\\", "\\\\")))
        os.replace(tmp, path)
    return path


# ---------------------------------------------------------------- checks

def check_pq_map(t):
    code, out = run("pq_map.py", t["pq"])
    ck("pq_map: exit 0, two queries, one loaded", code == 0 and "2 queries, 1 loaded to a sheet" in out, out[:300])
    ck("pq_map: a visible sheet gets no state tag", "[visible]" not in out, out[:300])
    blocks = {b.split("\n", 1)[0]: b for b in re.split(r"\n(?=\[\d+\] )", out)[1:]}
    raw, net = blocks.get("[1] SalesRaw", ""), blocks.get("[2] SalesNet", "")
    ck("pq_map: load target (sheet, table, range, columns)", "loads to : Data!SalesRaw A1:E6 | 4 query columns" in raw, raw)
    ck("pq_map: excel-side column", "excel-side columns (not from the query): Extra" in raw, raw)
    ck("pq_map: source path and picked sheet", "Source.xlsx" in raw and 'Sheet "Data"' in raw, raw)
    ck("pq_map: uses / used by", "used by  : SalesNet" in raw and "uses     : SalesRaw" in net, raw + net)
    ck("pq_map: connection-only query", "nothing on a sheet" in net, net)
    ck("pq_map: last fill", "LastUpdated=2026-01-02T03:04:05.0000000Z, Count=5, Status=Complete" in raw, raw)
    code, out = run("pq_map.py", t["pq"], "--query", "SalesNet")
    ck("pq_map: --query prints one query", code == 0 and "shared SalesNet" in out and "Table.SelectRows" in out and "shared SalesRaw" not in out, out)
    exp = Path(t["dir"]) / "export"
    code, out = run("pq_map.py", t["pq"], "--export", exp)
    files = sorted(p.name for p in exp.glob("*.pq")) if exp.is_dir() else []
    starts = all((exp / f).read_text(encoding="utf-8").startswith("section") for f in files)
    ck("pq_map: --export writes two .pq files", code == 0 and len(files) == 2 and starts, f"{files} {out}")
    code, out = run("pq_map.py", t["pq"], "--json")
    try:
        q = json.loads(out)[0]["queries"]
        good = [x["name"] for x in q] == ["SalesRaw", "SalesNet"] and q[0]["loads_to"][0]["excel_side_columns"] == ["Extra"]
    except (ValueError, KeyError, IndexError):
        good = False
    ck("pq_map: --json parses", code == 0 and good, out[:300])
    code, out = run("pq_map.py", t["base"])
    ck("pq_map: workbook without Power Query", code == 0 and "no Power Query" in out, out)
    code, out = run("xlsx_check.py", t["pq"])
    ck("xlsx_check: Power Query fixture is structurally ok", code == 0 and "structure ok" in out, out)


def check_xlsxpatch(t):
    d, base = Path(t["dir"]), t["base"]
    spec = {"set_value": [{"sheet": "Data", "cell": "B2", "value": 40}, {"sheet": "Data", "cell": "A3", "value": "Zed"},
                          {"sheet": "Data", "cell": "I2", "value": True}],
            "set_formula": [{"sheet": "Data", "cell": "G5", "formula": "=MAX(C2:C6)"}],
            "rewrite": [{"sheet": "Data", "cells": "E2:E6", "old": "=C{r}-D{r}", "new": "=C{r}-D{r}-1"},
                        {"sheet": "Data", "cells": "J2:J6", "old": "C{r}*2", "new": "C{r}*3"}],
            "table_rewrite": [{"table": "Sales", "column": "Net", "old": "C{r}-D{r}", "new": "C{r}-D{r}-1"}]}
    out = d / "patched.xlsx"
    code, log = run("xlsxpatch.py", "apply", base, out, jwrite(d / "spec.json", spec))
    ck("xlsxpatch apply: exit 0", code == 0 and out.is_file(), log)
    t["patched"] = out
    if code == 0:
        ws = load_workbook(out)["Data"]
        ck("xlsxpatch: set_value number, text, true", (ws["B2"].value, ws["A3"].value, ws["I2"].value) == (40, "Zed", True),
           str((ws["B2"].value, ws["A3"].value, ws["I2"].value)))
        ck("xlsxpatch: set_formula", ws["G5"].value == "=MAX(C2:C6)", str(ws["G5"].value))
        ck("xlsxpatch: rewrite with and without leading =", ws["E3"].value == "=C3-D3-1" and ws["J4"].value == "=C4*3",
           f"{ws['E3'].value} {ws['J4'].value}")
        ck("xlsxpatch: table_rewrite", "<calculatedColumnFormula>C2-D2-1<" in part(out, "xl/tables/table1.xml"), part(out, "xl/tables/table1.xml"))
    code, log = run("xlsxpatch.py", "cached", base, d / "cached.xlsx", jwrite(d / "vals.json", {"Data": {"E2": 7.5}}))
    ck("xlsxpatch cached: stores a result", code == 0 and load_workbook(d / "cached.xlsx", data_only=True)["Data"]["E2"].value == 7.5, log)
    t["cached"] = d / "cached.xlsx"

    bad = jwrite(d / "bad.json", {"rewrite": [{"sheet": "Data", "cells": "E2:E6", "old": "=C{r}+D{r}", "new": "=1"}]})
    code, log = run("xlsxpatch.py", "apply", base, d / "bad.xlsx", bad)
    ck("xlsxpatch refuses: wrong old formula", code == 1 and "mismatch" in log and not (d / "bad.xlsx").exists(), f"{code} {log}")
    code, log = run("xlsxpatch.py", "apply", base, base, jwrite(d / "empty.json", {}))
    ck("xlsxpatch refuses: SRC equals DST", code == 1 and "same file" in log, f"{code} {log}")
    code, log = run("xlsxpatch.py", "apply", base, d / "bad2.xlsx",
                    jwrite(d / "s.json", {"set_value": [{"sheet": "Nope", "cell": "A1", "value": 1}]}))
    ck("xlsxpatch refuses: unknown sheet, lists sheets", code == 1 and "Data" in log and "Summary" in log, f"{code} {log}")
    code, log = run("xlsxpatch.py", "apply", base, d / "bad3.xlsx",
                    jwrite(d / "s2.json", {"set_value": [{"sheet": "Data", "cell": "E2", "value": 1}]}))
    ck("xlsxpatch refuses: set_value on a formula cell", code == 1 and "formula" in log and not (d / "bad3.xlsx").exists(), f"{code} {log}")


def check_xlsx_check(t):
    d = Path(t["dir"])
    for label, path in (("base", t["base"]), ("patched output", t["patched"]), ("cached output", t["cached"]),
                        ("shared-formula fixture", t["shared"])):
        code, out = run("xlsx_check.py", path)
        ck(f"xlsx_check ok: {label}", code == 0 and "structure ok" in out, out)
    from openpyxl import Workbook
    wb = Workbook()
    wb.active["A1"], wb.active["B1"] = 2, "=A1*2"
    wb.save(d / "raw_openpyxl.xlsx")                    # openpyxl stores formulas with an empty <v></v>
    code, out = run("xlsx_check.py", d / "raw_openpyxl.xlsx")
    ck("xlsx_check ok: workbook straight from openpyxl", code == 0 and "structure ok" in out and "no stored result" in out, out)

    def broken(label, src, edit, expect):
        dst = d / f"broken_{len(label)}_{label[:6].replace(' ', '_')}.xlsx"
        zip_edit(src, dst, edit)
        code, out = run("xlsx_check.py", dst)
        ck(f"xlsx_check FAIL: {label}", code == 1 and any(l.startswith("FAIL") for l in out.splitlines()) and expect in out, f"{code} {out}")
    broken("two cells swapped in a row", t["base"],
           sub_part("xl/worksheets/sheet1.xml", r'(<c r="A2".*?</c>)(<c r="B2".*?</c>)', r"\2\1", re.S), "out of order")
    broken("table header differs from column name", t["base"], sub_part("xl/tables/table1.xml", r'name="Net"', 'name="NetX"'), "header")

    def cut(p):
        p["xl/worksheets/sheet2.xml"] = p["xl/worksheets/sheet2.xml"][:60]
    broken("truncated XML part", t["base"], cut, "not well-formed")
    broken("shared dependent without master", t["shared"],
           sub_part("xl/worksheets/sheet1.xml", r'<f t="shared" ref="E2:E6" si="0">', "<f>"), "without a master")
    junk = d / "junk.xlsx"
    junk.write_text("this is not a zip file", encoding="utf-8")
    code, out = run("xlsx_check.py", junk)
    ck("xlsx_check: non-zip gives FAIL, no traceback", code == 1 and "FAIL" in out and "Traceback" not in out, f"{code} {out}")


def check_xlsx_diff(t):
    d = Path(t["dir"])
    spec = {"set_value": [{"sheet": "Data", "cell": "B2", "value": 40}], "set_formula": [{"sheet": "Data", "cell": "I3", "formula": "=B3*2"}]}
    new = d / "diffed.xlsx"
    run("xlsxpatch.py", "apply", t["base"], new, jwrite(d / "dspec.json", spec))
    code, out = run("xlsx_diff.py", t["base"], new)
    lines = [l for l in out.splitlines() if l.startswith("('")]
    ck("xlsx_diff: names exactly the changed cells", code == 0 and len(lines) == 2 and "B2" in lines[0] and "I3" in lines[1], out)
    ck("xlsx_diff: count line", "2 cells differ in 1 sheets" in out, out)


def check_colinsert(t):
    d = Path(t["dir"])

    def insert(src, spec, name):
        dst = d / name
        code, out = run("colinsert.py", src, dst, jwrite(d / (name + ".json"), spec))
        return code, out, dst

    col = {"name": "Margin", "formula": "Sales[[#This Row],[Net]]/Sales[[#This Row],[Gross]]"}
    code, out, dst = insert(t["base"], {"sheet": "Data", "at": "F", "header_row": 1, "first_row": 2, "last_row": 6, "columns": [col]}, "append.xlsx")
    tab = part(dst, "xl/tables/table1.xml") if dst.exists() else ""
    ws = load_workbook(dst)["Data"] if dst.exists() else None
    chk = run("xlsx_check.py", dst)
    ck("colinsert append: table ref and column count grow", code == 0 and 'ref="A1:F6"' in tab and 'tableColumns count="6"' in tab, f"{code} {out} {tab}")
    ck("colinsert append: header and formula", ws is not None and ws["F1"].value == "Margin" and ws["F2"].value == "=" + col["formula"],
       str(ws and (ws["F1"].value, ws["F2"].value)))
    ck("colinsert append: xlsx_check ok", chk[0] == 0 and "structure ok" in chk[1], chk[1])

    t["mid"] = build_base(d / "mid_src.xlsx", summary_col="B")
    code, out, dst = insert(t["mid"], {"sheet": "Data", "at": "C", "header_row": 1, "first_row": 2, "last_row": 6, "columns": [{"name": "Spare"}]}, "middle.xlsx")
    ws = load_workbook(dst)["Data"] if dst.exists() else None
    chk = run("xlsx_check.py", dst)
    ck("colinsert middle: xlsx_check ok", code == 0 and chk[0] == 0 and "structure ok" in chk[1], f"{code} {out} {chk[1]}")
    got = ws and (ws["C1"].value, ws["D1"].value, ws["F2"].value, ws["H2"].value, ws["K2"].value)
    ck("colinsert middle: formulas to the right shift", got == ("Spare", "Gross", "=D2-E2", "=SUM(F2:F6)", "=D2*2"), str(got))

    code, out, dst = insert(t["base"], {"sheet": "Data", "at": "C", "header_row": 1, "first_row": 2, "last_row": 6, "columns": [{"name": "Spare"}]}, "guard.xlsx")
    ck("colinsert guard: Summary points at a moved column", code == 1 and "Summary" in out and not dst.exists(), f"{code} {out}")

    t["shared_mid"] = build_base(d / "shared_mid.xlsx", summary_col="B", shared=True)
    code, out, dst = insert(t["shared_mid"], {"sheet": "Data", "at": "C", "header_row": 1, "first_row": 2, "last_row": 6, "columns": [{"name": "Spare"}]}, "shared_ins.xlsx")
    xml = part(dst, "xl/worksheets/sheet1.xml") if dst.exists() else ""
    chk = run("xlsx_check.py", dst)
    ck("colinsert shared formulas: xlsx_check ok", code == 0 and chk[0] == 0 and "structure ok" in chk[1], f"{code} {out} {chk[1]}")
    ck("colinsert shared formulas: no escaped markup, master shifted",
       "&lt;" not in xml and '<f t="shared" ref="F2:F6" si="0">D2-E2</f>' in xml, xml[xml.find("<row r=\"2\""):][:600])


def check_xlsx_replace(t):
    d = Path(t["dir"]) / "replace"
    (d / "live").mkdir(parents=True)
    live, new, backups = d / "live" / "Book.xlsx", d / "new.xlsx", d / "backups"
    shutil.copyfile(t["base"], live)
    shutil.copyfile(t["patched"], new)
    old = live.read_bytes()
    sha = __import__("hashlib").sha256(old).hexdigest()[:16]
    common = [new, live, "--backup-dir", backups, "--wait", "0"]
    code, out = run("xlsx_replace.py", *common, "--base-sha", "0" * 16, "--go")
    ck("xlsx_replace: wrong base sha exits 2, live unchanged", code == 2 and live.read_bytes() == old and "changed since the base" in out, f"{code} {out}")
    lock = d / "live" / "~$Book.xlsx"
    lock.write_bytes(b"lock")
    code, out = run("xlsx_replace.py", *common, "--base-sha", sha, "--go")
    ck("xlsx_replace: lock file exits 2, live unchanged", code == 2 and live.read_bytes() == old and "lock file" in out, f"{code} {out}")
    lock.unlink()
    code, out = run("xlsx_replace.py", *common, "--base-sha", sha)
    ck("xlsx_replace: dry run exits 0, changes nothing", code == 0 and live.read_bytes() == old and "dry run" in out and not backups.exists(), f"{code} {out}")
    code, out = run("xlsx_replace.py", *common, "--base-sha", sha, "--go")
    saved = list(backups.glob("*_Book.xlsx")) if backups.is_dir() else []
    ck("xlsx_replace --go: live equals new", code == 0 and live.read_bytes() == new.read_bytes(), f"{code} {out}")
    ck("xlsx_replace --go: backup equals old", len(saved) == 1 and saved[0].read_bytes() == old, str(saved))
    ck("xlsx_replace --go: no .tmp file left", not any(p.name.endswith(".tmp") for p in live.parent.iterdir()), str(list(live.parent.iterdir())))


def find_lo():
    if os.environ.get("SMOKE_NO_LO"):
        return None
    for c in ("soffice", "libreoffice"):
        if shutil.which(c):
            return c
    mac = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    return mac if os.path.isfile(mac) else None


def check_lo_recalc(t):
    if not find_lo():
        print("skip lo_recalc (no LibreOffice)")
        return
    d = Path(t["dir"])
    vals = {"E%d" % (i + 2): g - c for i, (_, _, g, c) in enumerate(ROWS)}
    vals.update({"G2": sum(g - c for _, _, g, c in ROWS), "G3": sum(r[2] for r in ROWS) / len(ROWS)})
    vals.update({"J%d" % (i + 2): g * 2 for i, (_, _, g, _) in enumerate(ROWS)})
    spec = {"Data": vals, "Summary": {"B1": sum(r[2] for r in ROWS)}}
    good = d / "stored_ok.xlsx"
    run("xlsxpatch.py", "cached", t["base"], good, jwrite(d / "good.json", spec))
    code, out = run("lo_recalc.py", good)
    ck("lo_recalc: right stored results exit 0", code == 0 and "0 stored results differ" in out, f"{code} {out}")
    spec["Data"]["E3"] = 999
    wrong = d / "stored_bad.xlsx"
    run("xlsxpatch.py", "cached", t["base"], wrong, jwrite(d / "wrong.json", spec))
    code, out = run("lo_recalc.py", wrong)
    ck("lo_recalc: a wrong stored number exits 1 and names the cell", code == 1 and "E3" in out and "DIFFERENT" in out, f"{code} {out}")


def check_usage(t):
    for name in sorted(PY):
        code, out = run(name)
        ck(f"usage: {name} without arguments exits 2", code == 2 and name in out, f"{code} {out[:200]}")
    sh = PQ / "refresh_excel_mac.sh"
    if os.name == "nt" or not shutil.which("bash"):
        print("skip bash -n refresh_excel_mac.sh (no bash)")
        return
    p = subprocess.run(["bash", "-n", str(sh)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
    ck("refresh_excel_mac.sh: bash -n", p.returncode == 0, p.stdout.decode("utf-8", "replace"))


def main():
    t0 = time.time()
    with tempfile.TemporaryDirectory(prefix="smoke_") as tmp:
        d = Path(tmp)
        t = {"dir": d, "base": build_base(d / "base.xlsx"), "shared": build_base(d / "shared.xlsx", shared=True),
             "pq": build_pq(d / "pq.xlsx", d)}
        for fn in (check_pq_map, check_xlsxpatch, check_xlsx_check, check_xlsx_diff, check_colinsert,
                   check_xlsx_replace, check_lo_recalc, check_usage):
            try:
                fn(t)
            except Exception:
                RES["fail"] += 1
                print(f"FAIL {fn.__name__}: crashed\n" + traceback.format_exc())
    total = RES["ok"] + RES["fail"] + RES["known"]
    print(f"{RES['ok']} ok, {RES['fail']} failed, {RES['known']} known defects, {total} checks, {time.time() - t0:.1f}s")
    return 1 if RES["fail"] else 0


if __name__ == "__main__":
    sys.exit(main())
