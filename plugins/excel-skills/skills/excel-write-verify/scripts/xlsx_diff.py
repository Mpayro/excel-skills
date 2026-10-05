#!/usr/bin/env python3
"""Semantic diff of two workbooks: for every cell, what it holds (formula or constant) and its stored result.

  xlsx_diff.py BASE.xlsx NEW.xlsx

Prints the zip parts that were added, removed or changed, then one line per (sheet, column, kind)
with a count and two samples. Kinds:
  formula/const   what the cell holds changed
  cached          same formula, different stored result

Compares what cells hold and their stored results. It does not compare formats, merged cells, notes or
table definitions: the "parts changed" line shows that such a part changed.
Read-only. Needs openpyxl. Exit code 0 when it ran (the reader decides which differences were intended),
2 for a usage error or a file that cannot be read.
"""
import collections, sys, warnings, zipfile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula


def norm(x):
    if isinstance(x, ArrayFormula):
        return "{" + str(x.text) + "}"
    if isinstance(x, float):
        return round(x, 9)
    if hasattr(x, "__dict__"):          # other formula objects (data tables): compare by content
        return str(sorted(vars(x).items()))
    return x


def val(ws, rc):
    c = ws._cells.get(rc)
    return norm(c.value) if c is not None else None


def main(a, b):
    warnings.simplefilter("ignore")
    za, zb = zipfile.ZipFile(a), zipfile.ZipFile(b)
    na, nb = set(za.namelist()), set(zb.namelist())
    print("parts added:  ", sorted(nb - na) or "none")
    print("parts removed:", sorted(na - nb) or "none")
    print("parts changed:", sorted(n for n in na & nb if za.read(n) != zb.read(n)) or "none")

    fa, va = load_workbook(a, data_only=False), load_workbook(a, data_only=True)
    fb, vb = load_workbook(b, data_only=False), load_workbook(b, data_only=True)
    if fa.sheetnames != fb.sheetnames:
        print("sheets only in BASE:", [s for s in fa.sheetnames if s not in fb.sheetnames],
              "| only in NEW:", [s for s in fb.sheetnames if s not in fa.sheetnames])
    tot = collections.Counter()
    samples = collections.defaultdict(list)
    for s in fa.sheetnames:
        if s not in fb.sheetnames:
            continue
        A, B, AV, BV = fa[s], fb[s], va[s], vb[s]
        if A.dimensions != B.dimensions:
            print(f"{s}: used range {A.dimensions} -> {B.dimensions}")
        # Only cells that exist: a sheet can declare a used range of a million empty rows, and
        # A.cell(r, c) would create every one of them in memory.
        for r, c in sorted(set(A._cells) | set(B._cells)):
            x, y = val(A, (r, c)), val(B, (r, c))
            xv, yv = val(AV, (r, c)), val(BV, (r, c))
            kind = "formula/const" if x != y else "cached" if xv != yv else None
            if kind:
                key = (s, get_column_letter(c), kind)
                tot[key] += 1
                if len(samples[key]) < 2:
                    old, new = (x, y) if kind != "cached" else (xv, yv)
                    samples[key].append((f"{get_column_letter(c)}{r}", str(old)[:90], str(new)[:90]))
    for key, n in sorted(tot.items()):
        print(key, n, samples[key])
    print(f"{sum(tot.values())} cells differ in {len({k[0] for k in tot})} sheets")


if __name__ == "__main__":
    if len(sys.argv) == 2 and sys.argv[1] in ("-h", "--help"):
        print(__doc__)
        raise SystemExit(0)
    if len(sys.argv) != 3:
        print(__doc__, file=sys.stderr)
        raise SystemExit(2)
    try:
        main(sys.argv[1], sys.argv[2])
    except (OSError, zipfile.BadZipFile) as e:
        print(f"cannot read the workbook: {e}", file=sys.stderr)
        raise SystemExit(2) from None
